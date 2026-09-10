import json
import sqlite3
import openpyxl
import yaml
from pathlib import Path
from snipglide.database.connection import get_connection
from snipglide.database.snippet_repo import add_snippet, get_snippet_by_shortcut
from snipglide.database.group_repo import get_all_groups, add_group
from snipglide.models.snippet import Snippet
from snipglide.models.group import Group
from snipglide.utils.logger import logger

MAX_IMPORT_FILE_BYTES = 25 * 1024 * 1024

def _is_reasonable_import_file(file_path: str) -> bool:
    try:
        return Path(file_path).stat().st_size <= MAX_IMPORT_FILE_BYTES
    except OSError:
        return False

def export_backup(file_path: str) -> bool:
    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            
            def _fetch_table(table_name: str) -> list[dict]:
                try:
                    cursor.execute(f"SELECT * FROM {table_name}")
                    return [dict(row) for row in cursor.fetchall()]
                except Exception:
                    return []

            groups = _fetch_table("groups")
            snippets = _fetch_table("snippets")
            autocorrect = _fetch_table("autocorrect")
            note_categories = _fetch_table("note_categories")
            notes = _fetch_table("notes")
            note_settings = _fetch_table("note_settings")
            chat_note_sections = _fetch_table("chat_note_sections")
            chat_notes = _fetch_table("chat_notes")
            screenshot_folders = _fetch_table("screenshot_folders")
            screenshots = _fetch_table("screenshots")
            saved_regexes = _fetch_table("saved_regexes")
            saved_api = _fetch_table("saved_api_requests")
            dev_projects = _fetch_table("developer_projects")
            terminal_cmds = _fetch_table("terminal_commands")
            usage_history = _fetch_table("usage_history")
            
            backup_data = {
                "version": "2.0",
                "groups": groups,
                "snippets": snippets,
                "autocorrect": autocorrect,
                "note_categories": note_categories,
                "notes": notes,
                "note_settings": note_settings,
                "chat_note_sections": chat_note_sections,
                "chat_notes": chat_notes,
                "screenshot_folders": screenshot_folders,
                "screenshots": screenshots,
                "saved_regexes": saved_regexes,
                "saved_api_requests": saved_api,
                "developer_projects": dev_projects,
                "terminal_commands": terminal_cmds,
                "usage_history": usage_history,
            }
            
            with open(file_path, "w", encoding="utf-8") as f:
                json.dump(backup_data, f, indent=4, ensure_ascii=False)
                
            logger.info(f"Backup exported successfully to {file_path}")
            return True
    except Exception as e:
        logger.error(f"Failed to export backup: {e}")
        return False

def import_backup(file_path: str) -> bool:
    try:
        if not _is_reasonable_import_file(file_path):
            logger.error("Backup import rejected: file is too large.")
            return False
        with open(file_path, "r", encoding="utf-8") as f:
            backup_data = json.load(f)
            
        if not isinstance(backup_data, dict) or "snippets" not in backup_data or "groups" not in backup_data:
            logger.error("Invalid backup file: missing snippets or groups schema.")
            return False

        # Create automatic pre-restore safety snapshot before modifying database
        try:
            from snipglide.core.config import BACKUP_DIR
            safety_file = BACKUP_DIR / "pre_restore_safety_backup.json"
            export_backup(str(safety_file))
            logger.info(f"Pre-restore safety snapshot created at {safety_file}")
        except Exception as e:
            logger.warning(f"Could not create pre-restore safety snapshot: {e}")
            
        # Collect API requests for post-commit secure processing
        _pending_api_requests = backup_data.get("saved_api_requests", []) if "saved_api_requests" in backup_data else []

        with get_connection() as conn:
            cursor = conn.cursor()
            
            # Atomic restore of core snippet tables
            cursor.execute("DELETE FROM snippets")
            cursor.execute("DELETE FROM groups")
            cursor.execute("DELETE FROM autocorrect")
            cursor.execute("DELETE FROM sqlite_sequence WHERE name IN ('snippets', 'groups', 'autocorrect')")
            
            for g in backup_data.get("groups", []):
                if not isinstance(g, dict):
                    continue
                cursor.execute("""
                    INSERT INTO groups (id, name, description, icon, color, is_collapsed) 
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (
                    g.get("id"), g.get("name"), g.get("description", ""),
                    g.get("icon", "G"), g.get("color", "#2563eb"), g.get("is_collapsed", 0)
                ))
                
            for s in backup_data.get("snippets", []):
                if not isinstance(s, dict):
                    continue
                cursor.execute("""
                    INSERT INTO snippets (
                        id, shortcut, replacement, group_id, tags, description, language,
                        enabled, favorite, usage_counter, created_date, modified_date,
                        hotkey, regex_enabled, app_filter, window_filter, notes, snippet_type
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    s.get("id"), s.get("shortcut"), s.get("replacement"), s.get("group_id"),
                    s.get("tags", ""), s.get("description", ""), s.get("language", "Plain Text"),
                    s.get("enabled", 1), s.get("favorite", 0), s.get("usage_counter", 0),
                    s.get("created_date"), s.get("modified_date"), s.get("hotkey", ""),
                    s.get("regex_enabled", 0), s.get("app_filter", ""), s.get("window_filter", ""),
                    s.get("notes", ""), s.get("snippet_type", "Text")
                ))
                
            for a in backup_data.get("autocorrect", []):
                if not isinstance(a, dict):
                    continue
                cursor.execute("""
                    INSERT OR REPLACE INTO autocorrect (typo, correction) 
                    VALUES (?, ?)
                """, (a.get("typo"), a.get("correction")))

            # Restore Note Categories & Notes if present
            if "note_categories" in backup_data and backup_data["note_categories"]:
                cursor.execute("DELETE FROM notes")
                cursor.execute("DELETE FROM note_categories")
                cursor.execute("DELETE FROM sqlite_sequence WHERE name IN ('notes', 'note_categories')")
                for nc in backup_data.get("note_categories", []):
                    if isinstance(nc, dict) and nc.get("name"):
                        cursor.execute("""
                            INSERT INTO note_categories (id, name, icon, color, description)
                            VALUES (?, ?, ?, ?, ?)
                        """, (nc.get("id"), nc.get("name"), nc.get("icon", "N"), nc.get("color", "#2563eb"), nc.get("description", "")))

                for n in backup_data.get("notes", []):
                    if isinstance(n, dict) and n.get("title"):
                        cursor.execute("""
                            INSERT INTO notes (id, title, content, category_id, created_date, modified_date, color, pinned)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        """, (
                            n.get("id"), n.get("title"), n.get("content", ""), n.get("category_id"),
                            n.get("created_date"), n.get("modified_date"), n.get("color", "#2563eb"), n.get("pinned", 0)
                        ))

            if "note_settings" in backup_data:
                for ns in backup_data.get("note_settings", []):
                    if isinstance(ns, dict) and ns.get("key"):
                        cursor.execute("INSERT OR REPLACE INTO note_settings (key, value) VALUES (?, ?)", (ns["key"], ns.get("value", "")))

            # Restore Chat Note Sections & Chat Notes if present
            if "chat_note_sections" in backup_data and backup_data["chat_note_sections"]:
                cursor.execute("DELETE FROM chat_notes")
                cursor.execute("DELETE FROM chat_note_sections")
                cursor.execute("DELETE FROM sqlite_sequence WHERE name IN ('chat_notes', 'chat_note_sections')")
                for cns in backup_data.get("chat_note_sections", []):
                    if isinstance(cns, dict) and cns.get("name"):
                        cursor.execute("""
                            INSERT INTO chat_note_sections (id, name, icon, color)
                            VALUES (?, ?, ?, ?)
                        """, (cns.get("id"), cns.get("name"), cns.get("icon", "💬"), cns.get("color", "#25D366")))

                for cn in backup_data.get("chat_notes", []):
                    if isinstance(cn, dict) and cn.get("content"):
                        cursor.execute("""
                            INSERT INTO chat_notes (id, content, created_at, is_starred, section_id, tags, color)
                            VALUES (?, ?, ?, ?, ?, ?, ?)
                        """, (
                            cn.get("id"), cn.get("content"), cn.get("created_at"), cn.get("is_starred", 0),
                            cn.get("section_id", 1), cn.get("tags", ""), cn.get("color", "#25D366")
                        ))

            # Restore Screenshot Folders & Screenshots if present
            if "screenshot_folders" in backup_data and backup_data["screenshot_folders"]:
                for sf in backup_data.get("screenshot_folders", []):
                    if isinstance(sf, dict) and sf.get("name"):
                        cursor.execute("SELECT id FROM screenshot_folders WHERE name = ?", (sf["name"],))
                        if not cursor.fetchone():
                            cursor.execute("INSERT INTO screenshot_folders (name, color) VALUES (?, ?)", (sf["name"], sf.get("color", "#3b82f6")))

            if "screenshots" in backup_data and backup_data["screenshots"]:
                for sc in backup_data.get("screenshots", []):
                    if isinstance(sc, dict) and sc.get("file_path"):
                        cursor.execute("SELECT id FROM screenshots WHERE file_path = ?", (sc["file_path"],))
                        if not cursor.fetchone():
                            cursor.execute("""
                                INSERT INTO screenshots (
                                    file_path, filename, capture_type, width, height, file_size,
                                    created_at, is_favorite, note, duration, thumbnail_path, folder
                                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                            """, (
                                sc["file_path"], sc.get("filename", ""), sc.get("capture_type", "full"),
                                sc.get("width", 0), sc.get("height", 0), sc.get("file_size", 0),
                                sc.get("created_at"), sc.get("is_favorite", 0), sc.get("note", ""),
                                sc.get("duration", 0.0), sc.get("thumbnail_path", ""), sc.get("folder", "العامة")
                            ))

            # Optional developer tools restoration (idempotent / non-destructive)
            if "saved_regexes" in backup_data:
                try:
                    for r in backup_data.get("saved_regexes", []):
                        if isinstance(r, dict) and r.get("name") and r.get("pattern"):
                            cursor.execute("SELECT id FROM saved_regexes WHERE name = ?", (r["name"],))
                            if not cursor.fetchone():
                                cursor.execute(
                                    "INSERT INTO saved_regexes (name, pattern, description, flags, replacement, favorite) VALUES (?, ?, ?, ?, ?, ?)",
                                    (r.get("name"), r.get("pattern"), r.get("description", ""), r.get("flags", ""), r.get("replacement", ""), r.get("favorite", 0)),
                                )
                except Exception:
                    pass

            # API requests: determine which ones are NOT already in DB (deduplicate check)
            if _pending_api_requests:
                _to_insert_api = []
                try:
                    for req in _pending_api_requests:
                        if isinstance(req, dict) and req.get("name") and req.get("url"):
                            cursor.execute("SELECT id FROM saved_api_requests WHERE name = ? AND url = ?", (req["name"], req["url"]))
                            if not cursor.fetchone():
                                _to_insert_api.append(req)
                except Exception:
                    pass

            if "developer_projects" in backup_data:
                try:
                    for p in backup_data.get("developer_projects", []):
                        if isinstance(p, dict) and p.get("name") and p.get("project_path"):
                            cursor.execute("SELECT id FROM developer_projects WHERE project_path = ?", (p["project_path"],))
                            if not cursor.fetchone():
                                cursor.execute(
                                    "INSERT INTO developer_projects (name, project_path, language, framework, description, is_favorite) VALUES (?, ?, ?, ?, ?, ?)",
                                    (p["name"], p["project_path"], p.get("language", ""), p.get("framework", ""), p.get("description", ""), p.get("is_favorite", 0)),
                                )
                except Exception:
                    pass

            if "terminal_commands" in backup_data:
                try:
                    for c in backup_data.get("terminal_commands", []):
                        if isinstance(c, dict) and c.get("name") and c.get("command"):
                            cursor.execute("SELECT id FROM terminal_commands WHERE name = ? AND command = ?", (c["name"], c["command"]))
                            if not cursor.fetchone():
                                cursor.execute(
                                    "INSERT INTO terminal_commands (name, command, description, category, is_favorite) VALUES (?, ?, ?, ?, ?)",
                                    (c["name"], c["command"], c.get("description", ""), c.get("category", "General"), c.get("is_favorite", 0)),
                                )
                except Exception:
                    pass
                
            conn.commit()

        # After main connection is committed and closed, safely process API requests via
        # ApiRepository (which manages its own connection) to ensure encryption is applied.
        if _pending_api_requests and _to_insert_api:
            import json as _json
            from snipglide.database.api_repo import ApiRepository
            from snipglide.models.api_request import ApiRequest
            for req in _to_insert_api:
                try:
                    raw_auth = _json.loads(req.get("auth_data_json", "{}")) if isinstance(req.get("auth_data_json"), str) else (req.get("auth_data_json") or {})
                    raw_headers = _json.loads(req.get("headers_json", "[]")) if isinstance(req.get("headers_json"), str) else (req.get("headers_json") or [])
                    params = _json.loads(req.get("params_json", "[]")) if isinstance(req.get("params_json"), str) else (req.get("params_json") or [])
                    api_req = ApiRequest(
                        name=req["name"],
                        method=req.get("method", "GET"),
                        url=req["url"],
                        params=params,
                        headers=raw_headers,
                        auth_type=req.get("auth_type", "none"),
                        auth_data=raw_auth,
                        body_type=req.get("body_type", "none"),
                        body_content=req.get("body_content", ""),
                        collection_name=req.get("collection_name", "General"),
                        is_favorite=bool(req.get("is_favorite", 0)),
                    )
                    ApiRepository.add_request(api_req)
                except Exception as _e:
                    logger.error(f"Failed to add saved API request: {_e}")

        logger.info("Backup imported successfully.")
        return True

    except Exception as e:
        logger.error(f"Failed to import backup: {e}")
        return False

# excel export/import
def export_to_excel(snippets: list, file_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Snippets"
    ws.append(["Shortcut", "Replacement", "Group", "Tags", "Description", "Language", "Enabled", "Favorite"])
    
    groups = {g.id: g.name for g in get_all_groups()}
    for s in snippets:
        g_name = groups.get(s.group_id, "General")
        ws.append([
            s.shortcut, s.replacement, g_name, ",".join(s.tags),
            s.description, s.language, str(s.enabled), str(s.favorite)
        ])
    wb.save(file_path)

def import_from_excel(file_path: str) -> int:
    if not _is_reasonable_import_file(file_path):
        return 0
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    count = 0
    
    groups = {g.name.lower(): g.id for g in get_all_groups()}
    try:
        rows = ws.iter_rows(values_only=True)
        next(rows, None)
        for r in rows:
            if not r or len(r) < 2 or not r[0] or not r[1]:
                continue
            shortcut, replacement = str(r[0]), str(r[1])
            group_name = r[2] if len(r) > 2 else "General"
            tags_str = r[3] if len(r) > 3 else ""
            description = r[4] if len(r) > 4 else ""
            language = r[5] if len(r) > 5 else "Plain Text"
            enabled = r[6] if len(r) > 6 else True
            favorite = r[7] if len(r) > 7 else False
            
            g_name = str(group_name or "General")
            if g_name.lower() not in groups:
                g_id = add_group(Group(name=g_name, description="Imported via Excel"))
                groups[g_name.lower()] = g_id
            else:
                g_id = groups[g_name.lower()]
                
            tags = [t.strip() for t in str(tags_str or "").split(",") if t.strip()]
            existing = get_snippet_by_shortcut(shortcut)
            if existing:
                continue
                
            snippet = Snippet(
                shortcut=shortcut,
                replacement=replacement,
                group_id=g_id,
                tags=tags,
                description=str(description or ""),
                language=str(language or "Plain Text"),
                enabled=(str(enabled).lower() == "true" or enabled == 1 or enabled is True),
                favorite=(str(favorite).lower() == "true" or favorite == 1 or favorite is True)
            )
            add_snippet(snippet)
            count += 1
    finally:
        wb.close()
    return count

# yaml export/import
def export_to_yaml(snippets: list, file_path: str):
    groups = {g.id: g.name for g in get_all_groups()}
    data = []
    for s in snippets:
        data.append({
            "shortcut": s.shortcut,
            "replacement": s.replacement,
            "group": groups.get(s.group_id, "General"),
            "tags": s.tags,
            "description": s.description,
            "language": s.language,
            "enabled": s.enabled,
            "favorite": s.favorite
        })
    with open(file_path, "w", encoding="utf-8") as f:
        yaml.dump(data, f, allow_unicode=True)

def import_from_yaml(file_path: str) -> int:
    if not _is_reasonable_import_file(file_path):
        return 0
    with open(file_path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    if not isinstance(data, list):
        return 0
        
    groups = {g.name.lower(): g.id for g in get_all_groups()}
    count = 0
    for item in data:
        if not isinstance(item, dict):
            continue
        shortcut = item.get("shortcut")
        replacement = item.get("replacement")
        if not shortcut or not replacement:
            continue
            
        g_name = str(item.get("group", "General") or "General")
        if g_name.lower() not in groups:
            g_id = add_group(Group(name=g_name, description="Imported via YAML"))
            groups[g_name.lower()] = g_id
        else:
            g_id = groups[g_name.lower()]
            
        existing = get_snippet_by_shortcut(shortcut)
        if existing:
            continue
            
        snippet = Snippet(
            shortcut=shortcut,
            replacement=replacement,
            group_id=g_id,
            tags=item.get("tags", []),
            description=item.get("description", ""),
            language=item.get("language", "Plain Text"),
            enabled=bool(item.get("enabled", True)),
            favorite=bool(item.get("favorite", False))
        )
        add_snippet(snippet)
        count += 1
    return count

# json export/import
def export_to_json(snippets: list, file_path: str):
    groups = {g.id: g.name for g in get_all_groups()}
    data = []
    for s in snippets:
        data.append({
            "shortcut": s.shortcut,
            "replacement": s.replacement,
            "group": groups.get(s.group_id, "General"),
            "tags": s.tags,
            "description": s.description,
            "language": s.language,
            "enabled": s.enabled,
            "favorite": s.favorite
        })
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4, ensure_ascii=False)

def import_from_json(file_path: str) -> int:
    if not _is_reasonable_import_file(file_path):
        return 0
    with open(file_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, list):
        return 0
        
    groups = {g.name.lower(): g.id for g in get_all_groups()}
    count = 0
    for item in data:
        if not isinstance(item, dict):
            continue
        shortcut = item.get("shortcut")
        replacement = item.get("replacement")
        if not shortcut or not replacement:
            continue
            
        g_name = str(item.get("group", "General") or "General")
        if g_name.lower() not in groups:
            g_id = add_group(Group(name=g_name, description="Imported via JSON"))
            groups[g_name.lower()] = g_id
        else:
            g_id = groups[g_name.lower()]
            
        existing = get_snippet_by_shortcut(shortcut)
        if existing:
            continue
            
        snippet = Snippet(
            shortcut=shortcut,
            replacement=replacement,
            group_id=g_id,
            tags=item.get("tags", []),
            description=item.get("description", ""),
            language=item.get("language", "Plain Text"),
            enabled=bool(item.get("enabled", True)),
            favorite=bool(item.get("favorite", False))
        )
        add_snippet(snippet)
        count += 1
    return count
