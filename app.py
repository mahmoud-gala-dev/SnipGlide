import json
import os
import sys
import threading
import time
from pathlib import Path
from typing import Dict, List, Optional
from tkinter import filedialog

import customtkinter as ctk
from PIL import Image, ImageDraw
import pystray
from pynput import keyboard


APP_NAME = "SnipGlide Python"
DATA_DIR = Path(os.getenv("APPDATA", Path.home())) / "SnipGlidePython"
DATA_FILE = DATA_DIR / "snippets.json"
SETTINGS_FILE = DATA_DIR / "settings.json"

ctk.set_appearance_mode("system")
ctk.set_default_color_theme("blue")


def resource_path(relative: str) -> Path:
    if hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS) / relative
    return Path(__file__).resolve().parent / relative


class SnippetStore:
    def __init__(self):
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self.snippets: List[Dict] = []
        self.settings = {
            "enabled": True,
            "start_minimized": False,
            "case_sensitive": True,
            "max_buffer": 250,
        }
        self.load()

    def load(self):
        with self._lock:
            if DATA_FILE.exists():
                try:
                    self.snippets = json.loads(DATA_FILE.read_text(encoding="utf-8"))
                except Exception:
                    self.snippets = []
            else:
                # Try loading from sample_snippets.json
                sample_file = resource_path("sample_snippets.json")
                if sample_file.exists():
                    try:
                        self.snippets = json.loads(sample_file.read_text(encoding="utf-8"))
                    except Exception:
                        self.snippets = []
                else:
                    self.snippets = []

            # Backfill group key
            for item in self.snippets:
                if "group" not in item:
                    item["group"] = "General"

            if SETTINGS_FILE.exists():
                try:
                    loaded = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
                    self.settings.update(loaded)
                except Exception:
                    pass

    def save(self):
        with self._lock:
            DATA_FILE.write_text(
                json.dumps(self.snippets, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            SETTINGS_FILE.write_text(
                json.dumps(self.settings, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )

    def add(self, trigger: str, replacement: str, group: str = "General", enabled: bool = True):
        trigger = trigger.strip()
        if not trigger:
            raise ValueError("Shortcut trigger cannot be empty.")
        with self._lock:
            if any(item["trigger"] == trigger for item in self.snippets):
                raise ValueError("This shortcut already exists.")
            self.snippets.append({
                "trigger": trigger,
                "replacement": replacement,
                "group": group,
                "enabled": enabled,
            })
            self.save()

    def update(self, index: int, trigger: str, replacement: str, group: str, enabled: bool):
        trigger = trigger.strip()
        if not trigger:
            raise ValueError("Shortcut trigger cannot be empty.")
        with self._lock:
            for i, item in enumerate(self.snippets):
                if i != index and item["trigger"] == trigger:
                    raise ValueError("This shortcut already exists.")
            self.snippets[index] = {
                "trigger": trigger,
                "replacement": replacement,
                "group": group,
                "enabled": enabled,
            }
            self.save()

    def delete(self, index: int):
        with self._lock:
            del self.snippets[index]
            self.save()

    def active_map(self) -> Dict[str, str]:
        with self._lock:
            return {
                item["trigger"]: item["replacement"]
                for item in self.snippets
                if item.get("enabled", True)
            }


class ExpansionEngine:
    def __init__(self, store: SnippetStore, status_callback=None):
        self.store = store
        self.status_callback = status_callback
        self.buffer = ""
        self.controller = keyboard.Controller()
        self.listener: Optional[keyboard.Listener] = None
        self.running = False
        self.suspended = False
        self._lock = threading.RLock()

    def start(self):
        if self.listener is not None:
            return
        self.running = True
        self.listener = keyboard.Listener(on_press=self._on_press)
        self.listener.daemon = True
        self.listener.start()
        self._notify()

    def stop(self):
        self.running = False
        if self.listener:
            self.listener.stop()
            self.listener = None
        self._notify()

    def set_suspended(self, value: bool):
        self.suspended = value
        self.buffer = ""
        self._notify()

    def _notify(self):
        if self.status_callback:
            try:
                self.status_callback(self.running and not self.suspended)
            except Exception:
                pass

    def _on_press(self, key):
        if not self.running or self.suspended or not self.store.settings.get("enabled", True):
            return

        try:
            if key == keyboard.Key.backspace:
                self.buffer = self.buffer[:-1]
                return

            if key in {
                keyboard.Key.space, keyboard.Key.enter, keyboard.Key.tab,
                keyboard.Key.esc, keyboard.Key.left, keyboard.Key.right,
                keyboard.Key.up, keyboard.Key.down, keyboard.Key.home,
                keyboard.Key.end, keyboard.Key.page_up, keyboard.Key.page_down,
                keyboard.Key.delete,
            }:
                self.buffer = ""
                return

            char = getattr(key, "char", None)
            if char is None:
                return

            self.buffer += char
            max_len = int(self.store.settings.get("max_buffer", 250))
            self.buffer = self.buffer[-max_len:]

            snippets = self.store.active_map()
            case_sensitive = self.store.settings.get("case_sensitive", True)

            matched_trigger = None
            replacement = None

            for trigger, repl in sorted(snippets.items(), key=lambda x: len(x[0]), reverse=True):
                if case_sensitive:
                    matched = self.buffer.endswith(trigger)
                else:
                    matched = self.buffer.lower().endswith(trigger.lower())
                if matched:
                    matched_trigger = trigger
                    replacement = repl
                    break

            if matched_trigger is not None:
                self._expand(matched_trigger, replacement)

        except Exception:
            self.buffer = ""

    def _expand(self, trigger: str, replacement: str):
        with self._lock:
            self.suspended = True
            try:
                for _ in trigger:
                    self.controller.press(keyboard.Key.backspace)
                    self.controller.release(keyboard.Key.backspace)
                    time.sleep(0.002)

                self.controller.type(replacement)
                self.buffer = ""
            finally:
                time.sleep(0.03)
                self.suspended = False


class SnipGlideApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.store = SnippetStore()
        self.engine = ExpansionEngine(self.store, self._engine_status_changed)
        self.selected_index: Optional[int] = None
        self.tray_icon = None

        self.title(APP_NAME)
        self.geometry("980x650")
        self.minsize(850, 560)
        self.protocol("WM_DELETE_WINDOW", self.hide_window)

        self._build_ui()
        self._update_group_dropdowns()
        self._refresh_list()
        self.engine.start()
        self._start_tray()

        if self.store.settings.get("start_minimized", False):
            self.after(200, self.hide_window)

    def _build_ui(self):
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        sidebar = ctk.CTkFrame(self, width=250, corner_radius=0)
        sidebar.grid(row=0, column=0, sticky="nsew")
        sidebar.grid_propagate(False)

        ctk.CTkLabel(
            sidebar,
            text="SnipGlide",
            font=ctk.CTkFont(size=28, weight="bold"),
        ).pack(pady=(28, 4))

        ctk.CTkLabel(
            sidebar,
            text="Your shortcuts, your way",
            text_color=("gray35", "gray70"),
        ).pack(pady=(0, 22))

        self.status_label = ctk.CTkLabel(
            sidebar,
            text="● Running in background",
            text_color="#22c55e",
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.status_label.pack(pady=8)

        self.enable_switch = ctk.CTkSwitch(
            sidebar,
            text="Enable Replacement",
            command=self._toggle_enabled,
        )
        self.enable_switch.pack(pady=12)
        if self.store.settings.get("enabled", True):
            self.enable_switch.select()

        self.case_switch = ctk.CTkSwitch(
            sidebar,
            text="Case Sensitive",
            command=self._toggle_case,
        )
        self.case_switch.pack(pady=12)
        if self.store.settings.get("case_sensitive", True):
            self.case_switch.select()

        self.min_switch = ctk.CTkSwitch(
            sidebar,
            text="Start Minimized",
            command=self._toggle_start_minimized,
        )
        self.min_switch.pack(pady=12)
        if self.store.settings.get("start_minimized", False):
            self.min_switch.select()

        ctk.CTkButton(
            sidebar,
            text="Add New Shortcut",
            height=42,
            command=self._new_snippet,
        ).pack(fill="x", padx=24, pady=(28, 10))

        ctk.CTkButton(
            sidebar,
            text="Import from Excel",
            height=42,
            fg_color=("gray70", "gray30"),
            text_color=("black", "white"),
            command=self._import_excel,
        ).pack(fill="x", padx=24, pady=6)

        ctk.CTkButton(
            sidebar,
            text="Download Template",
            height=42,
            fg_color="transparent",
            border_width=1,
            command=self._download_template,
        ).pack(fill="x", padx=24, pady=6)

        ctk.CTkButton(
            sidebar,
            text="Hide to System Tray",
            fg_color="transparent",
            border_width=1,
            command=self.hide_window,
        ).pack(fill="x", padx=24, pady=(6, 10))

        ctk.CTkLabel(
            sidebar,
            text="Accepts any text or symbols:\n! @ # $ % ^ & * ( ) _ +",
            justify="center",
            text_color=("gray40", "gray65"),
        ).pack(side="bottom", pady=24)

        main = ctk.CTkFrame(self, corner_radius=0, fg_color=("gray96", "gray10"))
        main.grid(row=0, column=1, sticky="nsew")
        main.grid_columnconfigure(0, weight=1)
        main.grid_rowconfigure(1, weight=1)

        header = ctk.CTkFrame(main, fg_color="transparent")
        header.grid(row=0, column=0, sticky="ew", padx=28, pady=(24, 12))
        header.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            header,
            text="Manage Snippets",
            font=ctk.CTkFont(size=24, weight="bold"),
        ).grid(row=0, column=0, sticky="w")

        filter_frame = ctk.CTkFrame(header, fg_color="transparent")
        filter_frame.grid(row=0, column=1, sticky="e")

        self.group_filter = ctk.CTkOptionMenu(
            filter_frame,
            values=["All Groups"],
            command=lambda _: self._refresh_list(),
            width=140
        )
        self.group_filter.pack(side="left", padx=(0, 10))

        self.search_entry = ctk.CTkEntry(
            filter_frame,
            width=200,
            placeholder_text="Search shortcut...",
        )
        self.search_entry.pack(side="left")
        self.search_entry.bind("<KeyRelease>", lambda _e: self._refresh_list())

        content = ctk.CTkFrame(main)
        content.grid(row=1, column=0, sticky="nsew", padx=28, pady=(0, 28))
        content.grid_columnconfigure(0, weight=1)
        content.grid_columnconfigure(1, weight=2)
        content.grid_rowconfigure(0, weight=1)

        self.list_frame = ctk.CTkScrollableFrame(content, label_text="Snippets")
        self.list_frame.grid(row=0, column=0, sticky="nsew", padx=(14, 7), pady=14)

        editor = ctk.CTkFrame(content)
        editor.grid(row=0, column=1, sticky="nsew", padx=(7, 14), pady=14)
        editor.grid_columnconfigure(0, weight=1)
        editor.grid_rowconfigure(6, weight=1)

        ctk.CTkLabel(
            editor,
            text="Snippet Editor",
            font=ctk.CTkFont(size=19, weight="bold"),
        ).grid(row=0, column=0, sticky="w", padx=20, pady=(20, 12))

        ctk.CTkLabel(editor, text="Shortcut (Trigger)").grid(
            row=1, column=0, sticky="w", padx=20, pady=(4, 4)
        )
        self.trigger_entry = ctk.CTkEntry(
            editor,
            placeholder_text='e.g., #sig, $date, or !@#',
            height=40,
        )
        self.trigger_entry.grid(row=2, column=0, sticky="ew", padx=20)

        ctk.CTkLabel(editor, text="Group").grid(
            row=3, column=0, sticky="w", padx=20, pady=(12, 4)
        )
        group_frame = ctk.CTkFrame(editor, fg_color="transparent")
        group_frame.grid(row=4, column=0, sticky="ew", padx=20)
        group_frame.grid_columnconfigure(0, weight=1)

        self.group_var = ctk.StringVar(value="General")
        self.group_menu = ctk.CTkOptionMenu(
            group_frame,
            variable=self.group_var,
            values=["General"],
            height=40
        )
        self.group_menu.grid(row=0, column=0, sticky="ew", padx=(0, 10))

        self.add_group_btn = ctk.CTkButton(
            group_frame,
            text="+",
            width=40,
            height=40,
            command=self._add_new_group
        )
        self.add_group_btn.grid(row=0, column=1, sticky="e")

        ctk.CTkLabel(editor, text="Replacement Text").grid(
            row=5, column=0, sticky="w", padx=20, pady=(12, 4)
        )
        self.replacement_text = ctk.CTkTextbox(editor, wrap="word")
        self.replacement_text.grid(row=6, column=0, sticky="nsew", padx=20)

        self.item_enabled = ctk.CTkCheckBox(editor, text="Snippet Enabled")
        self.item_enabled.grid(row=7, column=0, sticky="w", padx=20, pady=14)
        self.item_enabled.select()

        buttons = ctk.CTkFrame(editor, fg_color="transparent")
        buttons.grid(row=8, column=0, sticky="ew", padx=20, pady=(0, 20))
        buttons.grid_columnconfigure((0, 1, 2), weight=1)

        ctk.CTkButton(
            buttons, text="Save", command=self._save_snippet
        ).grid(row=0, column=0, sticky="ew", padx=(0, 5))

        ctk.CTkButton(
            buttons,
            text="New",
            fg_color=("gray70", "gray30"),
            command=self._new_snippet,
        ).grid(row=0, column=1, sticky="ew", padx=5)

        ctk.CTkButton(
            buttons,
            text="Delete",
            fg_color="#dc2626",
            hover_color="#b91c1c",
            command=self._delete_snippet,
        ).grid(row=0, column=2, sticky="ew", padx=(5, 0))

    def _add_new_group(self):
        dialog = ctk.CTkInputDialog(text="Enter new group name:", title="Add Group")
        group_name = dialog.get_input()
        if group_name:
            group_name = group_name.strip()
            if group_name:
                existing_values = list(self.group_menu.cget("values"))
                if group_name not in existing_values:
                    existing_values.append(group_name)
                    sorted_groups = sorted(existing_values)
                    self.group_menu.configure(values=sorted_groups)
                self.group_var.set(group_name)

    def _update_group_dropdowns(self):
        existing_groups = set(item.get("group", "General") for item in self.store.snippets)
        existing_groups.add("General")
        sorted_groups = sorted(list(existing_groups))

        current_selection = self.group_var.get()
        self.group_menu.configure(values=sorted_groups)
        if current_selection not in sorted_groups:
            self.group_var.set("General")

        filter_groups = ["All Groups"] + sorted_groups
        current_filter = self.group_filter.get()
        self.group_filter.configure(values=filter_groups)
        if current_filter not in filter_groups:
            self.group_filter.set("All Groups")

    def _refresh_list(self):
        for widget in self.list_frame.winfo_children():
            widget.destroy()

        query = self.search_entry.get().strip().lower() if hasattr(self, "search_entry") else ""
        selected_filter = self.group_filter.get() if hasattr(self, "group_filter") else "All Groups"

        visible = []
        for i, item in enumerate(self.store.snippets):
            snippet_group = item.get("group", "General")
            if selected_filter != "All Groups" and snippet_group != selected_filter:
                continue

            text = f'{item["trigger"]} {item["replacement"]}'.lower()
            if not query or query in text:
                visible.append((i, item))

        if not visible:
            ctk.CTkLabel(
                self.list_frame,
                text="No snippets found.\nAdd your first snippet now.",
                text_color=("gray40", "gray65"),
            ).pack(pady=35)
            return

        for index, item in visible:
            preview = item["replacement"].replace("\n", " ")
            if len(preview) > 45:
                preview = preview[:45] + "…"

            button = ctk.CTkButton(
                self.list_frame,
                text=f'{item["trigger"]}\n{preview}',
                anchor="w",
                height=62,
                fg_color=("gray88", "gray18"),
                hover_color=("gray80", "gray25"),
                text_color=("gray10", "gray95"),
                command=lambda idx=index: self._select_snippet(idx),
            )
            button.pack(fill="x", pady=5)

    def _select_snippet(self, index: int):
        self.selected_index = index
        item = self.store.snippets[index]
        self.trigger_entry.delete(0, "end")
        self.trigger_entry.insert(0, item["trigger"])
        
        snippet_group = item.get("group", "General")
        self.group_var.set(snippet_group)
        
        self.replacement_text.delete("1.0", "end")
        self.replacement_text.insert("1.0", item["replacement"])
        if item.get("enabled", True):
            self.item_enabled.select()
        else:
            self.item_enabled.deselect()

    def _new_snippet(self):
        self.selected_index = None
        self.trigger_entry.delete(0, "end")
        self.group_var.set("General")
        self.replacement_text.delete("1.0", "end")
        self.item_enabled.select()
        self.trigger_entry.focus_set()

    def _save_snippet(self):
        trigger = self.trigger_entry.get()
        replacement = self.replacement_text.get("1.0", "end-1c")
        group = self.group_var.get()
        enabled = bool(self.item_enabled.get())

        try:
            if self.selected_index is None:
                self.store.add(trigger, replacement, group, enabled)
                self.selected_index = len(self.store.snippets) - 1
            else:
                self.store.update(self.selected_index, trigger, replacement, group, enabled)
            self._update_group_dropdowns()
            self._refresh_list()
            self._toast("Snippet saved successfully")
        except ValueError as exc:
            self._toast(str(exc), error=True)

    def _delete_snippet(self):
        if self.selected_index is None:
            self._toast("Select a snippet first", error=True)
            return
        self.store.delete(self.selected_index)
        self._new_snippet()
        self._update_group_dropdowns()
        self._refresh_list()
        self._toast("Snippet deleted")

    def _toggle_enabled(self):
        self.store.settings["enabled"] = bool(self.enable_switch.get())
        self.store.save()
        self.engine.set_suspended(not self.store.settings["enabled"])

    def _toggle_case(self):
        self.store.settings["case_sensitive"] = bool(self.case_switch.get())
        self.store.save()

    def _toggle_start_minimized(self):
        self.store.settings["start_minimized"] = bool(self.min_switch.get())
        self.store.save()

    def _engine_status_changed(self, active: bool):
        def update():
            if active and self.store.settings.get("enabled", True):
                self.status_label.configure(text="● Running in background", text_color="#22c55e")
            else:
                self.status_label.configure(text="● Suspended", text_color="#f59e0b")
        try:
            self.after(0, update)
        except Exception:
            pass

    def _toast(self, message: str, error: bool = False):
        popup = ctk.CTkToplevel(self)
        popup.overrideredirect(True)
        popup.attributes("-topmost", True)
        color = "#dc2626" if error else "#16a34a"
        label = ctk.CTkLabel(
            popup,
            text=message,
            fg_color=color,
            text_color="white",
            corner_radius=10,
            padx=18,
            pady=10,
        )
        label.pack()
        self.update_idletasks()
        x = self.winfo_x() + self.winfo_width() - 330
        y = self.winfo_y() + 35
        popup.geometry(f"+{x}+{y}")
        popup.after(1800, popup.destroy)

    def hide_window(self):
        self.withdraw()

    def show_window(self):
        self.after(0, self.deiconify)
        self.after(20, self.lift)
        self.after(30, self.focus_force)

    def _create_tray_image(self):
        image = Image.new("RGB", (64, 64), "white")
        draw = ImageDraw.Draw(image)
        draw.rounded_rectangle((6, 6, 58, 58), radius=15, fill=(37, 99, 235))
        draw.text((21, 14), "S", fill="white")
        return image

    def _start_tray(self):
        menu = pystray.Menu(
            pystray.MenuItem("Open SnipGlide", lambda: self.show_window(), default=True),
            pystray.MenuItem(
                "Enable / Disable",
                lambda: self.after(0, self._tray_toggle),
            ),
            pystray.MenuItem("Exit", lambda: self.after(0, self._quit_app)),
        )
        self.tray_icon = pystray.Icon(
            "SnipGlidePython",
            self._create_tray_image(),
            APP_NAME,
            menu,
        )
        threading.Thread(target=self.tray_icon.run, daemon=True).start()

    def _tray_toggle(self):
        if self.enable_switch.get():
            self.enable_switch.deselect()
        else:
            self.enable_switch.select()
        self._toggle_enabled()

    def _quit_app(self):
        self.engine.stop()
        if self.tray_icon:
            self.tray_icon.stop()
        self.destroy()

    def _import_excel(self):
        file_path = filedialog.askopenfilename(
            title="Import Excel File",
            filetypes=[("Excel Files", "*.xlsx")]
        )
        if not file_path:
            return

        try:
            import openpyxl
        except ImportError:
            self._toast("openpyxl library not installed.", error=True)
            return

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet = wb.active
            if sheet.max_row < 1:
                self._toast("Excel file is empty.", error=True)
                return

            # Read first row cells for headers
            headers = []
            for col in range(1, sheet.max_column + 1):
                val = sheet.cell(row=1, column=col).value
                headers.append(str(val).strip().lower() if val is not None else "")

            # Default column indices
            trigger_idx = 0
            replacement_idx = 1
            group_idx = 2
            enabled_idx = 3

            has_headers = False
            for i, h in enumerate(headers):
                if any(x in h for x in ["shortcut", "trigger", "الاختصار"]):
                    trigger_idx = i
                    has_headers = True
                elif any(x in h for x in ["replacement", "text", "البديل"]):
                    replacement_idx = i
                    has_headers = True
                elif any(x in h for x in ["group", "category", "المجموعة"]):
                    group_idx = i
                    has_headers = True
                elif any(x in h for x in ["enabled", "active", "مفعل"]):
                    enabled_idx = i
                    has_headers = True

            start_row = 2 if has_headers else 1

            imported_count = 0
            updated_count = 0

            # Read rows
            for row_num in range(start_row, sheet.max_row + 1):
                trigger_val = sheet.cell(row=row_num, column=trigger_idx + 1).value
                if trigger_val is None:
                    continue
                trigger = str(trigger_val).strip()
                if not trigger:
                    continue

                replacement_val = sheet.cell(row=row_num, column=replacement_idx + 1).value
                replacement = str(replacement_val) if replacement_val is not None else ""

                # Group
                group_val = sheet.cell(row=row_num, column=group_idx + 1).value if sheet.max_column > group_idx else None
                group = str(group_val).strip() if group_val is not None else "General"
                if not group:
                    group = "General"

                # Enabled
                enabled_val = sheet.cell(row=row_num, column=enabled_idx + 1).value if sheet.max_column > enabled_idx else None
                enabled = True
                if enabled_val is not None:
                    ev_str = str(enabled_val).strip().lower()
                    if ev_str in ["false", "0", "no", "لا", "off"]:
                        enabled = False

                # Check duplicate trigger
                existing_idx = None
                for idx, item in enumerate(self.store.snippets):
                    if item["trigger"] == trigger:
                        existing_idx = idx
                        break

                if existing_idx is not None:
                    self.store.snippets[existing_idx] = {
                        "trigger": trigger,
                        "replacement": replacement,
                        "group": group,
                        "enabled": enabled
                    }
                    updated_count += 1
                else:
                    self.store.snippets.append({
                        "trigger": trigger,
                        "replacement": replacement,
                        "group": group,
                        "enabled": enabled
                    })
                    imported_count += 1

            self.store.save()
            self._update_group_dropdowns()
            self._refresh_list()
            self._toast(f"Imported {imported_count} new, updated {updated_count} snippets")
        except Exception as exc:
            self._toast(f"Failed to import: {str(exc)}", error=True)

    def _download_template(self):
        file_path = filedialog.asksaveasfilename(
            title="Download Excel Template",
            defaultextension=".xlsx",
            initialfile="snipglide_template.xlsx",
            filetypes=[("Excel Files", "*.xlsx")]
        )
        if not file_path:
            return

        try:
            import openpyxl
        except ImportError:
            self._toast("openpyxl library not installed.", error=True)
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Template"

            # Write headers
            ws.append(["Shortcut", "Replacement", "Group", "Enabled"])
            # Write samples
            ws.append(["#sig", "Best regards,\nJohn Doe", "Work", "TRUE"])
            ws.append(["$date", "2026-07-24", "General", "TRUE"])
            ws.append(["!@#", "Custom keyboard shortcut snippet", "Personal", "TRUE"])

            wb.save(file_path)
            self._toast("Template saved successfully")
        except Exception as exc:
            self._toast(f"Failed to save template: {str(exc)}", error=True)


if __name__ == "__main__":
    app = SnipGlideApp()
    app.mainloop()
